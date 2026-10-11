"""Imports a completed vineyard onboarding template into the engine's data.

    python3 tools/import_vineyard.py path/to/completed.xlsx            # check only
    python3 tools/import_vineyard.py path/to/completed.xlsx --write    # check, then add
    python3 tools/import_vineyard.py path/to/completed.xlsx --write --replace
        (replace this vineyard's wines if it was imported before)
    python3 tools/import_vineyard.py --drop-dummy --write
        (remove the demo wines once real vineyards are loaded)

Checks every answer against the same lists the template's dropdowns use
(tools/onboarding_lists.py), so pasted or mistyped values are caught before
anything reaches the engine. Nothing is written unless the whole file passes.

Writes:
  api/data/wines.json       wines the match engine scores (no contact details)
  api/data/vineyards.json   vineyard profiles and stories, for "learn more" pages
  data/private/contacts.json contact details (git-ignored, never served)
"""
import json
import os
import re
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import onboarding_lists as L  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WINES = os.path.join(ROOT, "api", "data", "wines.json")
VINEYARDS = os.path.join(ROOT, "api", "data", "vineyards.json")
CONTACTS = os.path.join(ROOT, "data", "private", "contacts.json")
VINEYARD_MIN = L.MIN_ORDER  # bottles per vineyard per order, as in api/match.js

VD_FIELDS = {
    "Vineyard / winery name *": ("name", None, True),
    "Contact name *": ("contactName", None, True),
    "Contact email *": ("email", None, True),
    "Contact phone *": ("phone", None, True),
    "Website or Instagram": ("web", None, False),
    "State *": ("state", L.STATES, True),
    "Main region *": ("region", L.region_list(), True),
    "Ownership *": ("ownership", L.OWNERSHIP, True),
    "Years operating *": ("years", L.YEARS_OPERATING, True),
    "Typical annual production *": ("production", L.PRODUCTION, True),
    "Export experience *": ("exportExperience", L.EXPORT_EXPERIENCE, True),
    "Registered wine producer *": ("registered", L.REGISTERED, True),
    "Certifications": ("certification", L.CERTIFICATIONS, False),
    "Getting wine to Melbourne *": ("delivery", L.DELIVERY, True),
    f"Accept our {L.MIN_ORDER}-bottle minimum order *": ("acceptMin", L.ACCEPT_MIN_ORDER, True),
    "Cellar door": ("cellarDoor", L.YES_NO, False),
    "Comments about your vineyard": ("story", None, False),
}

WINE_COLS = [
    # (header prefix, key, allowed list or number spec, required)
    ("Wine name", "wineName", "text", True),
    ("Vintage", "vintage", L.VINTAGES, True),
    ("Style", "style", L.STYLES, True),
    ("Grape variety", "variety", L.VARIETIES, True),
    ("Region", "region", L.region_list(), True),
    ("Body", "body", L.BODY, True),
    ("Sweetness", "sweetness", L.SWEETNESS, True),
    ("Oak", "oak", L.OAK, False),
    ("Flavour note 1", "flavour1", L.FLAVOURS, True),
    ("Flavour note 2", "flavour2", L.FLAVOURS, False),
    ("Flavour note 3", "flavour3", L.FLAVOURS, False),
    ("Wholesale price", "price", ("decimal", 1, 2000), True),
    ("Bottles available", "stock", ("whole", 1, 100000), True),
    ("Bottle size", "bottleSize", L.BOTTLE_SIZES, True),
    ("Closure", "closure", L.CLOSURES, False),
    ("Alcohol", "alcohol", ("decimal", 5, 25), False),
    ("Exported before", "exportedBefore", L.EXPORTED_BEFORE, True),
    ("Highest award", "award", L.AWARDS, False),
    ("Comments about this wine", "notes", "text", False),
]


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def check_value(raw, spec, required, where, errors):
    val = clean(raw)
    if not val:
        if required:
            errors.append(f"{where}: required, but empty")
        return None
    if spec == "text":
        return val
    if isinstance(spec, tuple):
        kind, lo, hi = spec
        try:
            num = float(str(raw).replace("$", "").replace(",", ""))
        except ValueError:
            errors.append(f"{where}: should be a number, got \"{val}\"")
            return None
        if kind == "whole" and not num.is_integer():
            errors.append(f"{where}: should be a whole number, got {val}")
            return None
        if not lo <= num <= hi:
            errors.append(f"{where}: should be between {lo} and {hi}, got {val}")
            return None
        return int(num) if kind == "whole" else round(num, 2)
    # dropdown list: accept case differences, nothing else
    for option in spec:
        if option.lower() == val.lower():
            return option
    errors.append(f"{where}: \"{val}\" isn't one of the dropdown options")
    return None


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def read_template(path):
    wb = load_workbook(path, data_only=True)
    for sheet in ("Vineyard details", "Wines"):
        if sheet not in wb.sheetnames:
            raise SystemExit(f"This doesn't look like the Sapphire Panda template (no \"{sheet}\" tab).")
    errors = []

    vd = wb["Vineyard details"]
    vineyard = {}
    for row in vd.iter_rows(min_row=1, max_col=2):
        label = clean(row[0].value)
        if label in VD_FIELDS:
            key, options, required = VD_FIELDS[label]
            spec = options if options else "text"
            vineyard[key] = check_value(row[1].value, spec, required, f"Vineyard details, {label.replace(' *', '')}", errors)
    missing = [lbl for lbl in VD_FIELDS if VD_FIELDS[lbl][0] not in vineyard]
    if missing:
        errors.append("Vineyard details tab is missing rows: " + ", ".join(missing))
    if vineyard.get("acceptMin") == "No":
        errors.append(f"Vineyard details: doesn't accept the {L.MIN_ORDER}-bottle minimum order, so we can't feature this vineyard")
    email = vineyard.get("email")
    if email and not re.match(r"^[^\s@]+@[^\s@]+\.[^\s@]+$", email):
        errors.append(f"Vineyard details, Contact email: \"{email}\" doesn't look like an email address")

    ws = wb["Wines"]
    headers = [clean(c.value) for c in ws[1]]
    col_index = {}
    for prefix, key, _spec, _req in WINE_COLS:
        idx = next((i for i, h in enumerate(headers) if h.startswith(prefix)), None)
        if idx is None:
            errors.append(f"Wines tab: column \"{prefix}\" is missing")
        col_index[key] = idx

    wines = []
    seen = set()
    for r, row in enumerate(ws.iter_rows(min_row=3), start=3):
        values = [c.value for c in row]
        if not any(clean(v) for v in values):
            continue
        if clean(values[0]).lower().startswith("example:"):
            continue  # the template's sample row, if pasted into the input rows
        wine = {}
        for prefix, key, spec, required in WINE_COLS:
            idx = col_index.get(key)
            raw = values[idx] if idx is not None and idx < len(values) else None
            wine[key] = check_value(raw, spec, required, f"Wines row {r}, {prefix}", errors)
        ident = (str(wine.get("wineName")).lower(), wine.get("vintage"))
        if ident in seen:
            errors.append(f"Wines row {r}: {wine.get('wineName')} {wine.get('vintage')} is listed twice")
        seen.add(ident)
        wines.append(wine)
    if not wines:
        errors.append("Wines tab: no wines entered")
    return vineyard, wines, errors


def to_engine(vineyard, wines, start_id):
    out = []
    family = vineyard.get("ownership") == "Family-owned"
    scale = L.PRODUCTION_TO_SCALE.get(vineyard.get("production"), "")
    years = L.YEARS_MIDPOINT.get(vineyard.get("years"))
    for i, w in enumerate(wines):
        vintage = None if w["vintage"] == "NV" else int(w["vintage"])
        name = w["wineName"]
        if vintage and str(vintage) not in name:
            name = f"{name} {vintage}"
        elif not vintage and "NV" not in name:
            name = f"{name} NV"
        region = w["region"]
        state = vineyard["state"] if region in L.REGIONS.get(vineyard["state"], []) else (L.state_for_region(region) or vineyard["state"])
        flavours = [f for f in (w["flavour1"], w["flavour2"], w["flavour3"]) if f]
        out.append({
            "id": f"w{start_id + i}",
            "name": name,
            "vineyard": vineyard["name"],
            "state": state,
            "region": region if region != L.OTHER_REGION else vineyard["region"],
            "style": L.STYLE_TO_ENGINE.get(w["style"], w["style"]),
            "variety": w["variety"],
            "body": w["body"],
            "flavour": ", ".join(f.lower() for f in flavours),
            "vintage": vintage,
            "price": w["price"],
            "moq": VINEYARD_MIN,
            "stock": w["stock"],
            "vineyardScale": scale,
            "yearsOperating": years,
            "preferredMarket": "Other",
            "neverExported": w["exportedBefore"] == "Never exported",
            "familyOwned": family,
            "exportReady": True,
            # extra detail for future "learn more" views (ignored by scoring)
            "sweetness": w["sweetness"],
            "oak": w["oak"],
            "bottleSize": w["bottleSize"],
            "closure": w["closure"],
            "alcohol": w["alcohol"],
            "award": w["award"] if w["award"] != "None" else None,
            "notes": w["notes"],
            "source": "vineyard-template",
        })
    return out


def load(path, default):
    if os.path.exists(path):
        with open(path) as fh:
            return json.load(fh)
    return default


def save(path, data):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w") as fh:
        json.dump(data, fh, indent=2, ensure_ascii=False)
        fh.write("\n")


def main(argv):
    write = "--write" in argv
    replace = "--replace" in argv
    drop_dummy = "--drop-dummy" in argv
    files = [a for a in argv if not a.startswith("--")]

    wines_db = load(WINES, [])
    vineyards_db = load(VINEYARDS, [])
    contacts_db = load(CONTACTS, [])

    changed = False
    if drop_dummy:
        before = len(wines_db)
        wines_db = [w for w in wines_db if w.get("source") == "vineyard-template"]
        print(f"Demo wines removed: {before - len(wines_db)}")
        changed = before != len(wines_db)

    for path in files:
        vineyard, wines, errors = read_template(path)
        label = vineyard.get("name") or os.path.basename(path)
        if errors:
            print(f"\n{label}: {len(errors)} problem(s), nothing imported")
            for e in errors:
                print("  - " + e)
            continue
        existing = [w for w in wines_db if w.get("vineyard", "").lower() == vineyard["name"].lower()]
        if existing and not replace:
            print(f"\n{label}: already imported ({len(existing)} wines). Re-run with --replace to update it.")
            continue
        wines_db = [w for w in wines_db if w.get("vineyard", "").lower() != vineyard["name"].lower()]
        next_id = 1 + max([int(w["id"][1:]) for w in wines_db if re.match(r"^w\d+$", str(w.get("id")))] or [0])
        new = to_engine(vineyard, wines, next_id)
        wines_db.extend(new)
        vid = slug(vineyard["name"])
        profile = {k: vineyard.get(k) for k in ("name", "state", "region", "ownership", "years", "production",
                                                "exportExperience", "certification", "cellarDoor", "web", "story")}
        profile["id"] = vid
        vineyards_db = [v for v in vineyards_db if v.get("id") != vid] + [profile]
        contact = {"id": vid, "name": vineyard["name"], "contactName": vineyard["contactName"],
                   "email": vineyard["email"], "phone": vineyard["phone"], "delivery": vineyard["delivery"],
                   "registered": vineyard["registered"]}
        contacts_db = [c for c in contacts_db if c.get("id") != vid] + [contact]
        changed = True
        print(f"\n{label}: OK, {len(new)} wine(s) ready" + (" and added" if write else " (check only, use --write to add)"))
        for w in new:
            print(f"  + {w['name']}  |  {w['style']}, {w['variety']}, {w['region']}  |  ${w['price']:.2f}  |  {w['stock']} bottles")

    if write and changed:
        save(WINES, wines_db)
        save(VINEYARDS, vineyards_db)
        save(CONTACTS, contacts_db)
        print(f"\nSaved. The engine now has {len(wines_db)} wines.")


if __name__ == "__main__":
    main(sys.argv[1:])

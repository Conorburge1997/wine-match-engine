"""Builds the vineyard onboarding Excel template.

    python3 tools/build_vineyard_template.py [output.xlsx]

Every answer the engine relies on is a dropdown, so returned templates line
up exactly with the match algorithm. Free text is limited to names, contact
details and the comment fields. Numbers (price, bottles, alcohol) are
validated as numbers.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

sys.path.insert(0, __file__.rsplit("/", 1)[0])
import onboarding_lists as L  # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 else "sapphire-panda-vineyard-template.xlsx"
WINE_ROWS = 150  # data rows available on the Wines sheet

NAVY = "0B1430"
SAPPHIRE = "2F5BC0"
INPUT_FILL = PatternFill("solid", fgColor="EEF3FF")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
EXAMPLE_FILL = PatternFill("solid", fgColor="E4E7EE")
F = "Arial"
thin = Side(style="thin", color="C5CFE6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
OPEN = Protection(locked=False)
# Sheets are locked so headings, hints and instructions can't be edited or
# deleted; only the shaded answer cells are open. The password only stops
# accidental edits (it is not a security measure).
SHEET_PASSWORD = "sapphirepanda"


def lock(sheet):
    p = sheet.protection
    p.sheet = True
    p.password = SHEET_PASSWORD
    p.formatColumns = False   # still allow widening columns
    p.formatRows = False      # and resizing rows

wb = Workbook()

# ---------------- Lists (hidden) ----------------
lists = wb.active
lists.title = "Lists"
LISTS = {
    "Styles": L.STYLES, "Varieties": L.VARIETIES, "Regions": L.region_list(),
    "States": L.STATES, "Body": L.BODY, "Sweetness": L.SWEETNESS, "Oak": L.OAK,
    "Flavours": L.FLAVOURS, "Vintages": L.VINTAGES, "BottleSizes": L.BOTTLE_SIZES,
    "Closures": L.CLOSURES, "ExportedBefore": L.EXPORTED_BEFORE, "Awards": L.AWARDS,
    "Ownership": L.OWNERSHIP, "Years": L.YEARS_OPERATING, "Production": L.PRODUCTION,
    "ExportExp": L.EXPORT_EXPERIENCE, "Certs": L.CERTIFICATIONS, "Delivery": L.DELIVERY,
    "YesNo": L.YES_NO, "Registered": L.REGISTERED, "AcceptMin": L.ACCEPT_MIN_ORDER,
}
REF = {}
for col, (name, values) in enumerate(LISTS.items(), start=1):
    letter = get_column_letter(col)
    lists.cell(row=1, column=col, value=name).font = Font(name=F, bold=True)
    for i, v in enumerate(values, start=2):
        lists.cell(row=i, column=col, value=v).font = Font(name=F)
    REF[name] = f"=Lists!${letter}$2:${letter}${len(values) + 1}"
lists.sheet_state = "hidden"


def dropdown(ws, list_name, cell_range, prompt=None, required=False):
    dv = DataValidation(type="list", formula1=REF[list_name], allow_blank=not required)
    dv.promptTitle = "Tip"
    dv.error = "Please choose an option from the dropdown list."
    dv.errorTitle = "Choose from the list"
    dv.showErrorMessage = True
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(cell_range)


def number_rule(ws, cell_range, kind, lo, hi, message, prompt=None):
    dv = DataValidation(type=kind, operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True)
    if prompt:
        dv.promptTitle = "Tip"
        dv.prompt = prompt
        dv.showInputMessage = True
    dv.error = message
    dv.errorTitle = "Check this number"
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(cell_range)


# ---------------- Instructions ----------------
ins = wb.create_sheet("Instructions", 0)
ins.sheet_view.showGridLines = False
ins.column_dimensions["A"].width = 3
ins.column_dimensions["B"].width = 110
rows = [
    ("Sapphire Panda: vineyard onboarding template", Font(name=F, bold=True, size=16, color=NAVY)),
    ("Thank you for your interest in joining Sapphire Panda. This sheet tells us about your vineyard and the wines you would like to make available to overseas trade buyers.", None),
    ("", None),
    ("How to fill it in", Font(name=F, bold=True, size=12, color=SAPPHIRE)),
    ("1. Vineyard details tab: fill in the shaded cells about your vineyard (one sheet per vineyard).", None),
    ("2. Wines tab: add one row per wine you can make available. Use a new row for each vintage.", None),
    ("3. Most answers are dropdowns. Click a shaded cell and pick from the arrow. Typing your own answer will not be accepted.", None),
    ("4. Our minimum order is 12 bottles from each vineyard. Buyers can order as few as 12 bottles from you, and we only feature vineyards that accept this.", None),
    ("5. Prices are wholesale per 750 ml bottle in Australian dollars, ex-cellar, excluding GST and WET.", None),
    ("6. Use the comments columns for anything else: awards and where they were won, tasting notes, your history and story.", None),
    ("7. Save the file and email it back to us. Please also email a photo of each wine (bottle shot), named after the wine.", None),
    ("", None),
    ("Required fields are marked with *. If an option you need is missing, choose \"Other (see comments)\" and explain in the comments.", Font(name=F, italic=True, color="555555")),
    ("", None),
    ("Example", Font(name=F, bold=True, size=12, color=SAPPHIRE)),
    ("Row 2 of the Wines tab (grey) is a filled-in example of an ideal entry. It can't be edited and isn't imported. Start your own wines on row 3.", None),
    ("Click any cell on the Wines tab to see a short tip about what to enter.", None),
]
for i, (text, font) in enumerate(rows, start=1):
    c = ins.cell(row=i, column=2, value=text)
    c.font = font or Font(name=F, size=11)
    c.alignment = Alignment(wrap_text=True, vertical="top")

# ---------------- Vineyard details ----------------
vd = wb.create_sheet("Vineyard details", 1)
vd.sheet_view.showGridLines = False
vd.column_dimensions["A"].width = 38
vd.column_dimensions["B"].width = 46
vd.column_dimensions["C"].width = 60
vd["A1"] = "Vineyard details"
vd["A1"].font = Font(name=F, bold=True, size=14, color=NAVY)
vd["A2"] = "Fill in the shaded cells. Fields marked * are required."
vd["A2"].font = Font(name=F, italic=True, color="555555")
VD_FIELDS = [
    # (label, list name or None for free text, hint)
    ("Vineyard / winery name *", None, "As you would like it shown to buyers"),
    ("Contact name *", None, ""),
    ("Contact email *", None, ""),
    ("Contact phone *", None, ""),
    ("Website or Instagram", None, ""),
    ("State *", "States", ""),
    ("Main region *", "Regions", "Where your vineyard is based"),
    ("Ownership *", "Ownership", ""),
    ("Years operating *", "Years", ""),
    ("Typical annual production *", "Production", "All wines, in 9-litre cases"),
    ("Export experience *", "ExportExp", ""),
    ("Registered wine producer *", "Registered", ""),
    ("Certifications", "Certs", ""),
    ("Getting wine to Melbourne *", "Delivery", "We consolidate orders in Melbourne"),
    (f"Accept our {L.MIN_ORDER}-bottle minimum order *", "AcceptMin",
     f"Buyers can order as few as {L.MIN_ORDER} bottles from you. We only feature vineyards that accept this."),
    ("Cellar door", "YesNo", ""),
    ("Comments about your vineyard", None, "Your history and story, awards, what makes your wines special. Free text."),
]
VD_START = 4
for i, (label, list_name, hint) in enumerate(VD_FIELDS):
    r = VD_START + i
    a = vd.cell(row=r, column=1, value=label)
    a.font = Font(name=F, bold=True)
    a.alignment = Alignment(vertical="center")
    b = vd.cell(row=r, column=2)
    b.fill = INPUT_FILL
    b.border = BOX
    b.font = Font(name=F)
    b.alignment = Alignment(wrap_text=True, vertical="top")
    b.protection = OPEN
    vd.cell(row=r, column=3, value=hint).font = Font(name=F, italic=True, color="777777")
    vd.row_dimensions[r].height = 22
    if list_name:
        dropdown(vd, list_name, f"B{r}", required=label.endswith("*"))
last = VD_START + len(VD_FIELDS) - 1
vd.row_dimensions[last].height = 110

# ---------------- Wines ----------------
ws = wb.create_sheet("Wines", 2)
COLS = [
    # (header, width, list name / rule, required)
    ("Wine name *", 34, None, True),
    ("Vintage *", 11, "Vintages", True),
    ("Style *", 18, "Styles", True),
    ("Grape variety *", 30, "Varieties", True),
    ("Region *", 24, "Regions", True),
    ("Body *", 11, "Body", True),
    ("Sweetness *", 14, "Sweetness", True),
    ("Oak", 15, "Oak", False),
    ("Flavour note 1 *", 18, "Flavours", True),
    ("Flavour note 2", 18, "Flavours", False),
    ("Flavour note 3", 18, "Flavours", False),
    ("Wholesale price per bottle (AUD, ex-cellar, excl. GST & WET) *", 22, ("decimal", 1, 2000), True),
    ("Bottles available for export *", 16, ("whole", 1, 100000), True),
    ("Bottle size *", 16, "BottleSizes", True),
    ("Closure", 12, "Closures", False),
    ("Alcohol %", 11, ("decimal", 5, 25), False),
    ("Exported before? *", 18, "ExportedBefore", True),
    ("Highest award won", 20, "Awards", False),
    ("Comments about this wine", 60, None, False),
]
# Tips shown when a cell is selected (Excel's input message)
TIPS = [
    "The wine's name as on the label, without the vintage.",
    "Choose NV if non-vintage. Use a new row for each vintage.",
    "Choose from the list.",
    "Choose the closest. If it isn't listed, choose Other and explain in comments.",
    "Where the grapes were grown.",
    "Choose from the list.",
    "Choose from the list.",
    "Choose from the list.",
    "The main flavour. Up to three notes in total.",
    "Optional second flavour note.",
    "Optional third flavour note.",
    "Wholesale per bottle in AUD, ex-cellar, excluding GST and WET. Numbers only.",
    "How many bottles of this wine you could supply for export. Numbers only.",
    "Choose from the list.",
    "Choose from the list.",
    "e.g. 13.5",
    "Has this wine been sold outside Australia?",
    "The highest medal or trophy this wine has won. Give the show and year in comments.",
    "Awards and where they were won, how it's made, tasting notes, the story behind it.",
]
# Row 2: an ideal, fully filled-in entry (locked, not imported)
EXAMPLE = ["Example: Old Vine Shiraz", "2021", "Red", "Shiraz", "Barossa Valley", "Full", "Dry", "Oaked",
           "Blackberry", "Pepper", "Chocolate", 28.0, 600, "750 ml", "Screwcap", 14.5,
           "Never exported", "Silver medal",
           "Silver medal, Barossa Wine Show 2023. From dry-grown vines planted in 1960 on our home block. "
           "Hand-picked, open-fermented, 18 months in French oak. Drinking well now and will cellar to 2035."]
ws.freeze_panes = "B3"
ws.sheet_view.zoomScale = 100
for c, (head, width, rule, req) in enumerate(COLS, start=1):
    letter = get_column_letter(c)
    ws.column_dimensions[letter].width = width
    h = ws.cell(row=1, column=c, value=head)
    h.font = Font(name=F, bold=True, color="FFFFFF")
    h.fill = HEAD_FILL
    h.alignment = Alignment(wrap_text=True, vertical="center")
    h.border = BOX
    ex = ws.cell(row=2, column=c, value=EXAMPLE[c - 1])
    ex.font = Font(name=F, italic=True, color="4A5368")
    ex.fill = EXAMPLE_FILL
    ex.border = BOX
    ex.alignment = Alignment(wrap_text=True, vertical="top")
    rng = f"{letter}3:{letter}{WINE_ROWS + 2}"
    if isinstance(rule, str):
        dropdown(ws, rule, rng, prompt=TIPS[c - 1], required=req)
    elif isinstance(rule, tuple):
        kind, lo, hi = rule
        number_rule(ws, rng, kind, lo, hi, f"Please enter a number between {lo} and {hi}.", TIPS[c - 1])
    for r in range(3, WINE_ROWS + 3):
        cell = ws.cell(row=r, column=c)
        cell.fill = INPUT_FILL
        cell.border = BOX
        cell.font = Font(name=F)
        cell.protection = OPEN
    for r in range(2, WINE_ROWS + 3):  # example row gets the same formats
        cell = ws.cell(row=r, column=c)
        if c == 12:
            cell.number_format = '"$"#,##0.00'
        elif c == 13:
            cell.number_format = "#,##0"
        elif c == 16:
            cell.number_format = "0.0"
        if c == len(COLS):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[1].height = 58
ws.row_dimensions[2].height = 50
name_rule = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="80", allow_blank=True)
name_rule.error = "Please keep the wine name under 80 characters."
name_rule.showErrorMessage = True
name_rule.promptTitle = "Tip"
name_rule.prompt = TIPS[0]
name_rule.showInputMessage = True
ws.add_data_validation(name_rule)
name_rule.add(f"A3:A{WINE_ROWS + 2}")

comment_tip = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="2000", allow_blank=True)
comment_tip.promptTitle = "Tip"
comment_tip.prompt = TIPS[-1]
comment_tip.showInputMessage = True
ws.add_data_validation(comment_tip)
comment_tip.add(f"{get_column_letter(len(COLS))}3:{get_column_letter(len(COLS))}{WINE_ROWS + 2}")

for sheet in (ins, vd, ws, lists):
    lock(sheet)
wb.active = 0
wb.save(OUT)
print(OUT)

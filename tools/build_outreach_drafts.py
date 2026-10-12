"""Build the "Outreach drafts" sheet: one personalised first email per target winery.

Usage:
    python3 tools/build_outreach_drafts.py Vineyard_Database_v2.xlsx OUT.xlsx

Reads the "Vineyard Database" sheet (rows with Target 150 = Yes) and writes an
"Outreach drafts" sheet with subject, full email, what was personalised and what
to check before sending. Email copy agreed with Conor on 12 Oct 2026; see
claude/outreach-kit.md in the Sapphire Panda project. The workbook holds winery
contact details, so never commit it or the output.
"""
import re, sys, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)
ws = wb['Vineyard Database']
h = [c.value for c in ws[1]]
col = {k: h.index(k) for k in h if k}

GENERIC = {'Rosé', 'Rose', 'Sparkling', 'Unknown', 'Moscato', 'Fortified', 'Dessert', 'Red blend', 'White blend'}


def first_names(contact, winery):
    if not contact:
        return None
    seg = contact.split(';')[0]
    seg = re.sub(r'\(.*?\)', '', seg).strip()
    if not seg or 'family' in seg.lower():
        return None
    parts = [p.strip() for p in re.split(r'\s*&\s*|\s+and\s+', seg) if p.strip()]
    names = []
    for p in parts:
        tok = p.split()[0]
        if re.fullmatch(r"[A-Z][a-zA-Z'\-]+", tok) and tok.lower() not in ('the', 'mr', 'mrs', 'dr'):
            names.append(tok)
    if not names:
        return None
    return ' and '.join(names[:2])


def grapes(v):
    if not v or v == 'Unknown':
        return []
    gs = [g.strip() for g in str(v).split(',') if g.strip() and g.strip() not in GENERIC]
    return gs


SIGNATURE = {
    'Pyrenees': ['Shiraz'], 'Barossa Valley': ['Shiraz', 'Grenache'], 'Eden Valley': ['Riesling', 'Shiraz'],
    'Clare Valley': ['Riesling', 'Shiraz'], 'Coonawarra': ['Cabernet Sauvignon', 'Shiraz'], 'McLaren Vale': ['Shiraz', 'Grenache'],
    'Adelaide Hills': ['Chardonnay', 'Sauvignon Blanc', 'Pinot Noir'], 'Margaret River': ['Cabernet Sauvignon', 'Chardonnay'],
    'Great Southern': ['Riesling', 'Shiraz'], 'Yarra Valley': ['Pinot Noir', 'Chardonnay'], 'Mornington Peninsula': ['Pinot Noir', 'Chardonnay'],
    'Macedon Ranges': ['Pinot Noir', 'Chardonnay'], 'Geelong': ['Pinot Noir', 'Shiraz'], 'Heathcote': ['Shiraz'], 'Bendigo': ['Shiraz'],
    'Grampians': ['Shiraz'], 'Rutherglen': ['Durif', 'Shiraz'], 'Hunter Valley': ['Semillon', 'Shiraz'], 'Mudgee': ['Shiraz', 'Cabernet Sauvignon'],
    'Orange': ['Chardonnay', 'Shiraz'], 'Canberra': ['Shiraz', 'Riesling'], 'Hilltops': ['Shiraz'], 'Tumbarumba': ['Chardonnay', 'Pinot Noir'],
    'Granite Belt': ['Shiraz'], 'Southern Highlands': ['Pinot Noir', 'Chardonnay'], 'Swan Valley': ['Chenin Blanc', 'Verdelho', 'Shiraz'],
    'Geographe': ['Shiraz', 'Cabernet Sauvignon'], 'Gippsland': ['Pinot Noir', 'Chardonnay'], 'Henty': ['Riesling', 'Pinot Noir'],
}
PRIORITY = ['Shiraz', 'Pinot Noir', 'Chardonnay', 'Riesling', 'Cabernet Sauvignon', 'Grenache', 'Tempranillo', 'Semillon',
            'Sangiovese', 'Nebbiolo', 'Fiano', 'Vermentino', 'Pinot Gris', 'Durif', 'Mataro', 'Malbec', 'Merlot', 'Sauvignon Blanc']


def order_grapes(gs, region):
    sig = next((v for k, v in SIGNATURE.items() if k.lower() in (region or '').lower()), [])
    def key(g):
        if g in sig: return (0, sig.index(g))
        if g in PRIORITY: return (1, PRIORITY.index(g))
        return (2, 0)
    return sorted(gs, key=key)


def grape_phrase(gs):
    if not gs:
        return 'your wines'
    if len(gs) == 1:
        return 'your ' + gs[0]
    return 'your ' + gs[0] + ' and ' + gs[1]


def short(name):
    s = re.sub(r'\b(Wines?|Winery|Vineyards?|Estate Wines|Family Wines|Wine Co\.?)\b', '', name).strip(' .&-')
    return s if len(s) >= 3 else name


OPENERS = [
    "I came across {w} while looking at wineries in {r}, and {g} stood out.",
    "{w} came up when I was looking at {r} producers, and {g} caught my eye.",
    "I've been looking closely at wineries in {r}, and {w} is one I kept coming back to, particularly {g}.",
]
SUBJECT = "New overseas buyers for {s}?"


PARA1 = ("I'm getting in touch because I'm building a way for Australian wineries to sell overseas without the usual barriers: "
         "export compliance, minimum order quantities and developing international relationships. Most of Australia's wine exports "
         "go to a handful of big brands, largely because exporting is so costly and complicated for everyone else, and I think that "
         "can change. For you, it would simply mean new markets for your wine without extra time or complexity: when an order comes "
         "in, we collect the wine and take care of the rest.")

PARA2 = ("Sapphire Panda uses AI to match overseas buyers with Australian wineries on merit, not marketing budget. We build the "
         "relationships with buyers overseas, who find and order wines through our app, so all you need to do is have your wines "
         "listed. You give us your wholesale price per bottle, we set the overseas price, and we handle the export paperwork and "
         "freight. Orders start at one mixed case (12 bottles), shipped on a shared pallet with other wineries, and there's nothing "
         "to pay to be listed.")
# Softer alternative for the wholesale sentence, not yet confirmed by Conor:
# "You set the wholesale price you'd like per bottle, and we take care of everything from there: overseas pricing, export paperwork and freight."

CLOSE_VIC = ("I'm based in Melbourne, so I'm happy to visit the vineyard or start with a quick 15-minute call. Would either suit in "
             "the next couple of weeks? I'd love to work together and get your wine tasted abroad. If this isn't for you, no worries at all!")
CLOSE_INTERSTATE = ("I'm based in Melbourne and happy to start with a quick 15-minute call, and to visit when I'm next in {place}. "
                    "Would a call suit in the next couple of weeks? I'd love to work together and get your wine tasted abroad. "
                    "If this isn't for you, no worries at all!")

SIGN_OFF = """Cheers,
Conor Burge
Founder, Sapphire Panda
sapphirepanda.com | [mobile]

If you'd prefer not to receive emails from Sapphire Panda, just reply and let me know."""

FOLLOW_UP = """Hi {n},

Just following up on my email last week about Sapphire Panda. In short: we find overseas buyers for Australian wineries and handle the export side (compliance, paperwork, freight and the buyer relationships), so all you need to do is have your wines listed. Orders start at one mixed case, and there's nothing to pay to be listed.

Would a quick 15-minute call suit sometime in the next fortnight? If now isn't the right time, no worries at all.

Cheers,
Conor

If you'd prefer not to receive emails from Sapphire Panda, just reply and let me know."""

rows = []
i = 0
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[col['Target 150']] != 'Yes':
        continue
    w = r[col['Vineyard Name']]
    region = r[col['Region']] if r[col['Region']] not in (None, 'Unknown') else r[col['State']]
    names = first_names(r[col['Contact name']], w)
    hello = f"Hi {names}," if names else f"Hi {short(w)} team,"
    gs = order_grapes(grapes(r[col['Wine Varieties (findawinery.com blurb)']]), str(region))
    opener = OPENERS[i % 3].format(w=w, r=region, g=grape_phrase(gs))
    if r[col['State']] == 'VIC':
        close = CLOSE_VIC
    else:
        close = CLOSE_INTERSTATE.format(place=region)
    body = f"{hello}\n\n{opener} {PARA1}\n\n{PARA2}\n\n{close}\n\n{SIGN_OFF}"
    subj = SUBJECT.format(s=short(w))
    used = [f"Greeting: {('contact name ' + names) if names else 'no named contact, so addressed to the team'}",
            f"Region: {region}", f"Grapes: {', '.join(gs[:2]) if gs else 'not known, so kept general'}"]
    if r[col['Production (cases a year, Halliday)']]:
        used.append(f"Size: {r[col['Production (cases a year, Halliday)']]:,} cases (Halliday)")
    checks = ["Swap the opening line for something specific from their website (a wine, award or story) if you can."]
    email = r[col['Contact email']]
    if not email:
        checks.append("No email found: use their contact form or phone (see Outreach notes in the database).")
    if not gs:
        checks.append("Grapes unknown: name one of their wines in the opening line.")
    notes = r[col['Outreach notes']] or ''
    for flag in ('confirm it reaches', 'for sale', 'Only an events', 'Only a bookings', 'Several venue'):
        if flag in notes:
            checks.append(notes[notes.find(flag) - 40 if notes.find(flag) > 40 else 0:].split('|')[0].strip())
            break
    rows.append([w, r[col['State']], region, names or '', email or '', subj, body, '\n'.join(used), '\n'.join(checks), 'Not sent'])
    i += 1

if 'Outreach drafts' in wb.sheetnames:
    del wb['Outreach drafts']
od = wb.create_sheet('Outreach drafts', index=2)
head = ['Winery', 'State', 'Region', 'Greeting name', 'Email', 'Subject', 'Email body', 'What was personalised', 'Check before sending', 'Status']
od.append(head)
for row in rows:
    od.append(row)
hf = Font(name='Arial', bold=True, color='FFFFFF')
fill = PatternFill('solid', fgColor='1F3A6E')
for c in od[1]:
    c.font = hf; c.fill = fill; c.alignment = Alignment(vertical='center', wrap_text=True)
widths = [26, 7, 18, 16, 30, 40, 90, 34, 44, 12]
for idx, wdt in enumerate(widths):
    od.column_dimensions[openpyxl.utils.get_column_letter(idx + 1)].width = wdt
for row in od.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name='Arial', size=10)
        c.alignment = Alignment(vertical='top', wrap_text=True)
od.freeze_panes = 'B2'
od.auto_filter.ref = f"A1:J{len(rows) + 1}"
dv = DataValidation(type='list', formula1='"Not sent,Sent,Followed up,Replied - interested,Replied - not interested,No reply"', allow_blank=True)
od.add_data_validation(dv); dv.add(f"J2:J{len(rows) + 1}")

# follow-up template at the bottom
start = len(rows) + 3
od.cell(start, 1, 'Follow-up (send about a week later if no reply)').font = Font(name='Arial', bold=True)
od.cell(start + 1, 7, FOLLOW_UP.replace('{n}', '[first name]')).alignment = Alignment(vertical='top', wrap_text=True)
od.cell(start + 1, 7).font = Font(name='Arial', size=10)

for s in wb.worksheets:
    s.sheet_view.tabSelected = (s.title == wb.sheetnames[0])
wb.active = 0
wb.save(OUT)
print(len(rows), 'drafts written to', OUT)
